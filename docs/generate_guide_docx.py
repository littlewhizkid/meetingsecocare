from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

SCREENSHOTS = os.path.join(os.path.dirname(__file__), 'screenshots')

GREEN_DARK  = RGBColor(0x16, 0x65, 0x34)
GREEN_MID   = RGBColor(0x16, 0xa3, 0x4a)
GRAY_TEXT   = RGBColor(0x37, 0x41, 0x51)
AMBER_TEXT  = RGBColor(0x92, 0x40, 0x0e)
BLUE_TEXT   = RGBColor(0x1e, 0x40, 0xaf)
GRAY_LIGHT  = RGBColor(0x9c, 0xa3, 0xaf)

doc = Document()

# ── Page margins ────────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Cm(21)
section.page_height = Cm(29.7)
section.left_margin   = Cm(2.2)
section.right_margin  = Cm(2.2)
section.top_margin    = Cm(2)
section.bottom_margin = Cm(2)

# ── Default style ────────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.font.color.rgb = GRAY_TEXT

def add_heading(text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = GREEN_DARK
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after  = Pt(6)
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = GREEN_MID
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(4)
    return p

def add_body(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT
    p.paragraph_format.space_after = Pt(4)
    return p

def add_step(num, text):
    p = doc.add_paragraph(style='List Number')
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_note(text):
    p = doc.add_paragraph()
    run = p.add_run('Catatan: ' + text)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = AMBER_TEXT
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_after  = Pt(6)
    return p

def add_screenshot(name, caption=None):
    path = os.path.join(SCREENSHOTS, f'{name}.png')
    if not os.path.exists(path):
        return
    doc.add_picture(path, width=Cm(16.6))
    last_para = doc.paragraphs[-1]
    last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        run.font.size = Pt(8)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
        p.paragraph_format.space_after = Pt(8)

def add_divider():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '16a34a')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)


# ── Header ──────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p.add_run('ecoCare')
r1.bold = True
r1.font.size = Pt(22)
r1.font.color.rgb = GREEN_DARK
r2 = p.add_run(' Meeting Rooms')
r2.font.size = Pt(22)
r2.font.color.rgb = GREEN_MID
p.paragraph_format.space_after = Pt(2)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run('Panduan Pengguna — Pemesanan Ruang Meeting')
r.font.size = Pt(12)
r.font.color.rgb = GREEN_MID
p2.paragraph_format.space_after = Pt(2)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p3.add_run('https://meeting.ecocare.id')
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
p3.paragraph_format.space_after = Pt(8)

add_divider()

# ── Section 1: Masuk ─────────────────────────────────────────────────────────
add_heading('1.  Masuk ke Aplikasi')
add_body('Buka browser dan akses alamat berikut:')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('https://meeting.ecocare.id')
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = GREEN_DARK
p.paragraph_format.space_after = Pt(6)

add_step(1, 'Masukkan alamat email kerja Anda (contoh: namaanda@ecocare.co.id)')
add_step(2, 'Masukkan password Anda')
add_step(3, 'Klik tombol Sign In')
add_note('Jika Anda belum pernah login atau lupa password, hubungi administrator sistem untuk mereset password Anda.')
add_screenshot('signin', 'Tampilan halaman login ecoCare Meeting Rooms')

# ── Section 2: Kalender ──────────────────────────────────────────────────────
add_heading('2.  Tampilan Kalender Pemesanan')
add_body('Setelah masuk, Anda akan melihat kalender utama yang menampilkan keempat ruang meeting secara berdampingan untuk tanggal yang dipilih.')

# Room table
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ['Ruang', 'Kapasitas', 'Cocok untuk']
rows_data = [
    ('Board Room',        '12 orang', 'Rapat besar, presentasi'),
    ('Small Meeting Room','6 orang',  'Diskusi tim, pertemuan kecil'),
    ('Podcast Room',      '10 orang', 'Rekaman, sesi pelatihan'),
    ('Interview Room',    '4 orang',  'Wawancara, pertemuan pribadi'),
]
hdr_row = table.rows[0]
for i, h in enumerate(headers):
    cell = hdr_row.cells[i]
    cell.text = h
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    run.font.size = Pt(10)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '166534')
    tcPr.append(shd)

for i, (name, cap, use) in enumerate(rows_data):
    row = table.rows[i + 1]
    row.cells[0].text = name
    row.cells[1].text = cap
    row.cells[2].text = use
    fill = 'f0fdf4' if i % 2 == 0 else 'ffffff'
    for cell in row.cells:
        run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(cell.text)
        run.font.size = Pt(9)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

doc.add_paragraph()
add_body('Slot waktu yang kosong dapat diklik untuk melakukan pemesanan. Slot yang sudah terisi menampilkan detail booking yang ada.')
add_screenshot('calendar', 'Tampilan kalender pemesanan — empat ruang meeting ditampilkan berdampingan')

# ── Section 3: Navigasi Tanggal ──────────────────────────────────────────────
add_heading('3.  Navigasi Tanggal')
add_bullet('Klik panah kiri (←) untuk ke hari sebelumnya')
add_bullet('Klik panah kanan (→) untuk ke hari berikutnya')
add_bullet('Klik tanggal di tengah untuk membuka kalender dan langsung memilih tanggal tertentu')

# ── Section 4: Memesan Ruang ─────────────────────────────────────────────────
add_heading('4.  Cara Memesan Ruang Meeting')
add_step(1, 'Temukan ruang dan slot waktu yang Anda inginkan pada kalender')
add_step(2, 'Klik pada slot yang kosong — formulir pemesanan akan muncul')
add_step(3, 'Isi detail berikut:')

detail_table = doc.add_table(rows=4, cols=2)
detail_table.style = 'Table Grid'
details = [
    ('Waktu Mulai',  'Terisi otomatis dari slot yang diklik (dapat diubah)'),
    ('Waktu Selesai','Pilih waktu berakhirnya rapat'),
    ('Nama Anda',    'Terisi otomatis dari akun Anda'),
    ('Judul Rapat',  'Masukkan deskripsi singkat rapat'),
]
for i, (label, desc) in enumerate(details):
    row = detail_table.rows[i]
    row.cells[0].text = label
    row.cells[1].text = desc
    fill = 'f0fdf4' if i % 2 == 0 else 'ffffff'
    r0 = row.cells[0].paragraphs[0].runs[0]
    r0.bold = True
    r0.font.size = Pt(9)
    r1 = row.cells[1].paragraphs[0].runs[0]
    r1.font.size = Pt(9)
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

doc.add_paragraph()
add_step(4, 'Klik Konfirmasi Pemesanan')
add_step(5, 'Notifikasi konfirmasi berwarna hijau akan muncul di bagian atas layar')
add_note('Anda tidak dapat memesan slot waktu yang sudah terlewat.')

# ── Section 5: Lihat Pemesanan ───────────────────────────────────────────────
add_heading('5.  Melihat Pemesanan Anda')
add_step(1, 'Klik tombol My Bookings di pojok kanan atas layar')
add_step(2, 'Panel akan terbuka dan menampilkan semua pemesanan Anda yang akan datang')
add_step(3, 'Setiap pemesanan menampilkan nama ruang, tanggal, waktu, dan judul rapat')

# ── Section 6: Batalkan ──────────────────────────────────────────────────────
add_heading('6.  Membatalkan Pemesanan')
add_heading('Dari kalender:', level=2)
add_step(1, 'Temukan pemesanan Anda pada kalender (ditampilkan dalam warna hijau)')
add_step(2, 'Arahkan kursor ke atasnya — tombol X akan muncul di sudut kanan atas')
add_step(3, 'Klik X tersebut')
add_step(4, 'Konfirmasi pembatalan dengan mengklik Yes, Cancel Booking')

add_heading('Dari My Bookings:', level=2)
add_step(1, 'Buka My Bookings (tombol di pojok kanan atas)')
add_step(2, 'Temukan pemesanan yang ingin dibatalkan')
add_step(3, 'Klik Cancel di samping pemesanan tersebut')
add_step(4, 'Konfirmasi ketika diminta')
add_note('Anda hanya dapat membatalkan pemesanan milik Anda sendiri. Administrator dapat membatalkan pemesanan siapa pun.')

# ── Section 7: Keluar ────────────────────────────────────────────────────────
add_heading('7.  Keluar dari Aplikasi')
add_body('Klik ikon keluar (panah →) di samping nama Anda di pojok kanan atas layar.')

# ── Tips ─────────────────────────────────────────────────────────────────────
add_divider()
add_heading('Tips & Informasi Penting')
add_bullet('Bentrok jadwal? Sistem secara otomatis mencegah pemesanan yang tumpang tindih pada ruang yang sama.')
add_bullet('Butuh waktu lebih lama? Pilih waktu mulai terlebih dahulu, lalu pilih waktu selesai yang sesuai.')
add_bullet('Slot waktu lampau ditampilkan lebih redup dan tidak dapat diklik.')
add_bullet('Semua waktu menggunakan Waktu Indonesia Barat (WIB, UTC+7).')

# ── Bantuan ──────────────────────────────────────────────────────────────────
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Butuh bantuan? Hubungi administrator sistem atau tim IT Anda.')
r.font.size = Pt(9)
r.font.color.rgb = BLUE_TEXT

add_divider()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ecoCare Head Office  |  Sistem Pemesanan Ruang Meeting  |  meeting.ecocare.id')
r.font.size = Pt(8)
r.font.color.rgb = GRAY_LIGHT

# ── Save ─────────────────────────────────────────────────────────────────────
output_path = 'docs/Panduan-Pemesanan-Ruang-Meeting-ecoCare.docx'
doc.save(output_path)
print(f'Word document created: {output_path}')
